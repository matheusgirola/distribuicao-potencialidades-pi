# -*- coding: utf-8 -*-
"""Planilha de apoio dos dois relatorios (Estoque e CE+RIE): Leia-me, abas por
recorte, base municipal com filtro, e indicadores derivados por formula."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import comum as C

AZUL = Font(name='Arial', size=10, color='0000FF')     # dado extraído
PRETO = Font(name='Arial', size=10)                     # fórmula
NEG = Font(name='Arial', size=10, bold=True)
TIT = Font(name='Arial', size=12, bold=True, color='FFFFFF')
CINZA_IT = Font(name='Arial', size=9, italic=True, color='808080')
FILL_TIT = PatternFill('solid', fgColor='1F6F8B')
FILL_HDR = PatternFill('solid', fgColor='DCE6F1')
WRAP = Alignment(wrap_text=True, vertical='center', horizontal='center')
BORDA = Border(bottom=Side(style='thin', color='BFBFBF'))
FONTE_TXT = ('Fonte: elaboração própria a partir de IBGE (PAM, PPM, PEVS, Malha Municipal 2025) '
             'e RAIS/MTE (2025); matrizes de vizinhança DNIT/DER-PI e OpenRouteService. Extração em 06/08/2026.')

wb = Workbook(); wb.remove(wb.active)

def cab(ws, titulo, headers, widths):
    ws.sheet_view.showGridLines = False
    ncol = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c = ws.cell(1, 1, titulo); c.font = TIT; c.fill = FILL_TIT
    c.alignment = Alignment(vertical='center'); ws.row_dimensions[1].height = 22
    for j, h in enumerate(headers, 1):
        c = ws.cell(2, j, h); c.font = NEG; c.fill = FILL_HDR
        c.alignment = WRAP; c.border = BORDA
    ws.row_dimensions[2].height = 30
    for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    return 3   # primeira linha de dados

def rodape(ws, row, ncol):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    c = ws.cell(row, 1, FONTE_TXT); c.font = CINZA_IT

def num(ws, r, jc, v, fmt=None, font=AZUL, align=None):
    c = ws.cell(r, jc, v); c.font = font
    if fmt: c.number_format = fmt
    if align: c.alignment = Alignment(horizontal=align)
    return c

# ---------------- abas de subclasses ----------------
def aba_sub(nome, titulo, pkl):
    df = pd.read_pickle(pkl)
    ws = wb.create_sheet(nome)
    hdr = ['Subclasse','Municípios presentes','I de Moran global','p (999 perm.)',
           'Significativa (p<0,05)','Municípios sig. (999 p.)','Núcleos AA','Sig. nominais (9.999 p.)',
           'Pós-FDR (9.999 p.)','Robustos às 4 W alt.','Taxa de robustez (%)']
    r0 = cab(ws, titulo, hdr, [46,11,10,10,11,11,9,11,11,11,11])
    for i, row in df.iterrows():
        r = r0 + i
        num(ws, r, 1, row['subclasse'], font=AZUL)
        num(ws, r, 2, row['presentes'], '#,##0')
        num(ws, r, 3, row['I'], '0.0000')
        num(ws, r, 4, row['p'], '0.0000')
        num(ws, r, 5, row['sig'], '0')
        num(ws, r, 6, row['n_sig'], '#,##0')
        num(ws, r, 7, row['n_aa'], '#,##0')
        num(ws, r, 8, row['nom9999'], '#,##0')
        num(ws, r, 9, row['pos_fdr'], '#,##0')
        num(ws, r, 10, row['rob4'], '#,##0')
        num(ws, r, 11, f'=IF(F{r}=0,"–",J{r}/F{r}*100)', '0.0', PRETO)
    rN = r0 + len(df) - 1
    rT = rN + 1
    num(ws, rT, 1, 'Total', font=NEG)
    for col in ['F','G','H','I','J']:
        c = ws.cell(rT, ord(col)-64, f'=SUM({col}{r0}:{col}{rN})'); c.font = NEG; c.number_format='#,##0'
    num(ws, rT, 11, f'=J{rT}/F{rT}*100', '0.0', NEG)
    ws.freeze_panes = f'B{r0}'
    ws.auto_filter.ref = f'A2:K{rN}'
    rodape(ws, rT+2, 11)
    return r0, rN, rT

e_sub = aba_sub('1 Est Subclasses', 'Pipeline de estoques — LISA univariado por subclasse (T1+T2, base v2)',
                C.ROBUSTEZ / 'xls_uni_est.pkl')
s_sub = aba_sub('3 CERIE Subclasses', 'Pipeline shift-share (CE+RIE) — LISA univariado por subclasse (T1+T2)',
                C.ROBUSTEZ / 'xls_uni_ss.pkl')

# ---------------- abas de pares (núcleo robusto) ----------------
def aba_pares(nome, titulo, pkl):
    df = pd.read_pickle(pkl)
    ws = wb.create_sheet(nome)
    hdr = ['#','Subclasse focal (x)','Subclasse parceira (lag y)','I bivariado (Queen)',
           'p (9.999 perm.)','Núcleos AA','Copresentes']
    r0 = cab(ws, titulo, hdr, [5,46,46,12,11,9,11])
    for i, row in df.iterrows():
        r = r0 + i
        num(ws, r, 1, row['rank'], '0')
        num(ws, r, 2, row['focal']); num(ws, r, 3, row['parceira'])
        num(ws, r, 4, row['I'], '0.0000'); num(ws, r, 5, row['p'], '0.0000')
        num(ws, r, 6, row['n_aa'], '#,##0'); num(ws, r, 7, row['copresentes'], '#,##0')
    rN = r0 + len(df) - 1
    ws.freeze_panes = f'B{r0}'
    ws.auto_filter.ref = f'A2:G{rN}'
    rodape(ws, rN+2, 7)
    return r0, rN

e_par = aba_pares('2 Est Pares', 'Pipeline de estoques — núcleo robusto de pares (FDR e 4 matrizes alternativas)',
                  C.ROBUSTEZ / 'xls_core_est.pkl')
s_par = aba_pares('4 CERIE Pares', 'Pipeline shift-share (CE+RIE) — núcleo robusto de pares (FDR e 4 matrizes alternativas)',
                  C.ROBUSTEZ / 'xls_core_ss.pkl')

# ---------------- base municipal ----------------
me = pd.read_pickle(C.ROBUSTEZ / 'xls_mun_est.pkl').rename(columns={'aa_uni':'e_uni','aa_biv':'e_biv'})
ms = pd.read_pickle(C.ROBUSTEZ / 'xls_mun_ss.pkl').rename(columns={'aa_uni':'s_uni','aa_biv':'s_biv'})
mm = me.merge(ms, on='mun')
ws = wb.create_sheet('5 Base Municipal')
hdr = ['Município','Estoque: subclasses em núcleo AA','Estoque: pares robustos c/ núcleo AA',
       'CE+RIE: subclasses em núcleo AA','CE+RIE: pares robustos c/ núcleo AA','Total (soma)']
r0m = cab(ws, 'Base municipal — participação em núcleos Alto-Alto significativos (Queen, 999 perm.)',
          hdr, [28,15,15,15,15,10])
for i, row in mm.iterrows():
    r = r0m + i
    num(ws, r, 1, row['mun'])
    num(ws, r, 2, int(row['e_uni']), '#,##0'); num(ws, r, 3, int(row['e_biv']), '#,##0')
    num(ws, r, 4, int(row['s_uni']), '#,##0'); num(ws, r, 5, int(row['s_biv']), '#,##0')
    num(ws, r, 6, f'=SUM(B{r}:E{r})', '#,##0', PRETO)
rNm = r0m + len(mm) - 1
rTm = rNm + 1
num(ws, rTm, 1, 'Total', font=NEG)
for col in ['B','C','D','E','F']:
    c = ws.cell(rTm, ord(col)-64, f'=SUM({col}{r0m}:{col}{rNm})'); c.font = NEG; c.number_format='#,##0'
ws.freeze_panes = f'B{r0m}'
ws.auto_filter.ref = f'A2:F{rNm}'
rodape(ws, rTm+2, 6)

# ---------------- indicadores ----------------
ws = wb.create_sheet('6 Indicadores')
r0i = cab(ws, 'Indicadores derivados — calculados por fórmula sobre as abas anteriores',
          ['Indicador','Estoques','CE+RIE'], [62,14,14])
Q = lambda s: f"'{s}'!"
E, S = Q('1 Est Subclasses'), Q('3 CERIE Subclasses')
EP, SP = Q('2 Est Pares'), Q('4 CERIE Pares')
BM = Q('5 Base Municipal')
er0, ern, ert = e_sub; sr0, srn, srt = s_sub
epr0, eprn = e_par; spr0, sprn = s_par
linhas = [
 ('Subclasses elegíveis', f'=COUNT({E}C{er0}:C{ern})', f'=COUNT({S}C{sr0}:C{srn})'),
 ('Subclasses significativas (p<0,05)', f'=SUM({E}E{er0}:E{ern})', f'=SUM({S}E{sr0}:E{srn})'),
 ('Municípios sig. — total (999 perm.)', f'=SUM({E}F{er0}:F{ern})', f'=SUM({S}F{sr0}:F{srn})'),
 ('Municípios sig. — nominais (9.999 perm.)', f'=SUM({E}H{er0}:H{ern})', f'=SUM({S}H{sr0}:H{srn})'),
 ('Municípios sig. — pós-FDR (9.999 perm.)', f'=SUM({E}I{er0}:I{ern})', f'=SUM({S}I{sr0}:I{srn})'),
 ('Taxa de sobrevivência ao FDR (%)', '=C{r}/B{r}', '=D{r}/C{r}'),  # placeholder, montado abaixo
 ('Núcleos robustos às 4 W alternativas', f'=SUM({E}J{er0}:J{ern})', f'=SUM({S}J{sr0}:J{srn})'),
 ('Núcleo robusto de pares (contagem)', f'=COUNT({EP}A{epr0}:A{eprn})', f'=COUNT({SP}A{spr0}:A{sprn})'),
 ('I bivariado máximo do núcleo robusto', f'=MAX({EP}D{epr0}:D{eprn})', f'=MAX({SP}D{spr0}:D{sprn})'),
 ('I bivariado mediano do núcleo robusto', f'=MEDIAN({EP}D{epr0}:D{eprn})', f'=MEDIAN({SP}D{spr0}:D{sprn})'),
 ('Município líder em núcleos AA (univariado)',
    f'=INDEX({BM}A{r0m}:A{rNm},MATCH(MAX({BM}B{r0m}:B{rNm}),{BM}B{r0m}:B{rNm},0))',
    f'=INDEX({BM}A{r0m}:A{rNm},MATCH(MAX({BM}D{r0m}:D{rNm}),{BM}D{r0m}:D{rNm},0))'),
 ('Município líder em pares robustos (bivariado)',
    f'=INDEX({BM}A{r0m}:A{rNm},MATCH(MAX({BM}C{r0m}:C{rNm}),{BM}C{r0m}:C{rNm},0))',
    f'=INDEX({BM}A{r0m}:A{rNm},MATCH(MAX({BM}E{r0m}:E{rNm}),{BM}E{r0m}:E{rNm},0))'),
 ('Municípios com ao menos um núcleo AA (univariado)',
    f'=COUNTIF({BM}B{r0m}:B{rNm},">0")', f'=COUNTIF({BM}D{r0m}:D{rNm},">0")'),
]
for k,(rot,f1,f2) in enumerate(linhas):
    r = r0i + k
    num(ws, r, 1, rot, font=PRETO)
    if 'sobrevivência' in rot:
        rFDR = r-1; rTOT = r-2
        num(ws, r, 2, f'=B{rFDR}/B{rTOT}*100', '0.0', PRETO)
        num(ws, r, 3, f'=C{rFDR}/C{rTOT}*100', '0.0', PRETO)
    else:
        fmt = '0.0000' if 'I bivariado m' in rot else ('#,##0' if 'líder' not in rot else None)
        num(ws, r, 2, f1, fmt, PRETO); num(ws, r, 3, f2, fmt, PRETO)
rodape(ws, r0i+len(linhas)+1, 3)


# ---------------- duplo núcleo ----------------
dc = pd.read_pickle(C.ROBUSTEZ / 'xls_cross.pkl')
ws = wb.create_sheet('7 Duplo Nucleo')
hdr = ['#','Par de subclasses','I biv. Estoques','I biv. CE+RIE','Núcleos AA Estoques',
       'Núcleos AA CE+RIE','AA convergentes','Municípios convergentes']
r0d = cab(ws, 'Duplo núcleo — 32 pares presentes nos núcleos robustos dos dois pipelines',
          hdr, [5,62,12,12,12,12,12,70])
for i, row in dc.iterrows():
    r = r0d + i
    num(ws, r, 1, row['rank'], '0')
    num(ws, r, 2, row['par'])
    num(ws, r, 3, row['I_est'], '0.0000'); num(ws, r, 4, row['I_ss'], '0.0000')
    num(ws, r, 5, row['aa_est'], '#,##0'); num(ws, r, 6, row['aa_ss'], '#,##0')
    num(ws, r, 7, row['aa_conv'], '#,##0')
    num(ws, r, 8, row['municipios_conv'])
rNd = r0d + len(dc) - 1
ws.freeze_panes = f'C{r0d}'
ws.auto_filter.ref = f'A2:H{rNd}'
rodape(ws, rNd+2, 8)

# ---------------- Leia-me ----------------
ws = wb.create_sheet('Leia-me', 0)
ws.sheet_view.showGridLines = False
ws.column_dimensions['A'].width = 108
ws.merge_cells('A1:A1')
c = ws.cell(1,1,'Planilha de apoio — Análise de agrupamentos espaciais das potencialidades T1 e T2 — Piauí')
c.font = Font(name='Arial', size=13, bold=True, color='FFFFFF'); c.fill = FILL_TIT
ws.row_dimensions[1].height = 24
texto = [
 '',
 'OBJETIVO — Base dos dois relatórios analíticos: (i) pipeline de estoques (intensidade: estoque municipal no ano final, '
 'log1p; fontes PAM, PPM, PEVS e RAIS) e (ii) pipeline shift-share (intensidade: CE+RIE do emprego formal RAIS, '
 'log sinalizada). Recorte: classes de potencialidade T1 e T2, três janelas de comparação (bases 2013, 2018 e 2022), '
 'deduplicação pelo maior valor de intensidade por par município–subclasse.',
 '',
 'MÉTODO — LISA univariado (I de Moran local) e I de Moran bivariado, matriz Queen padronizada por linha, 999 '
 'permutações condicionais, semente 42, p < 0,05; subclasses elegíveis com presença em ≥ 10 municípios; pares com '
 '≥ 5 municípios copresentes. Robustez: (a) FDR de Benjamini–Hochberg com 9.999 permutações; (b) manutenção do '
 'resultado sob quatro matrizes alternativas de vizinhança (KNN-5, distância inversa euclidiana, distância rodoviária '
 'sobre a malha DNIT/DER-PI e tempo de viagem OpenRouteService/OpenStreetMap). O núcleo robusto de pares exige '
 'os dois critérios simultaneamente, com subclasse focal significativa e I bivariado positivo.',
 '',
 'RESSALVAS — Tempos de viagem modelados (não observados); quatro municípios com tempos imputados (R² = 0,933); '
 'no pipeline CE+RIE, Pau D\u2019Arco do Piauí e Pedro Laurentino não registram potencialidade classificada com '
 'decomposição válida. A coluna Pós-FDR conta municípios significativos em qualquer quadrante (inclui núcleos '
 'Baixo-Baixo de ausência conjunta) e usa 9.999 permutações, não sendo diretamente comparável às colunas de 999.',
 '',
 'CONVENÇÃO DE CORES — Azul: valor extraído dos pipelines do projeto (não editar). Preto: calculado por fórmula '
 'nesta planilha; recalcula ao alterar os dados.',
 '',
 'ÍNDICE — 1 Est Subclasses: LISA univariado por subclasse, pipeline de estoques (120 linhas). 2 Est Pares: núcleo '
 'robusto de pares do pipeline de estoques (340). 3 CERIE Subclasses: idem para o pipeline shift-share (87). '
 '4 CERIE Pares: núcleo robusto CE+RIE (50). 5 Base Municipal: os 224 municípios e suas contagens de núcleos AA. '
 '6 Indicadores: sínteses calculadas por fórmula. 7 Duplo Nucleo: os 32 pares presentes nos dois núcleos robustos, '
 'com os municípios em que os núcleos Alto-Alto dos dois pipelines coincidem.',
 '',
 FONTE_TXT,
]
for k, t in enumerate(texto, 2):
    c = ws.cell(k, 1, t); c.font = CINZA_IT if t==FONTE_TXT else Font(name='Arial', size=10)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if t and t != FONTE_TXT: ws.row_dimensions[k].height = max(15, 14*(len(t)//105+1))

wb.save(C.SAIDAS / 'Base_analise_espacial_potencialidades.xlsx')
print('planilha gravada')
